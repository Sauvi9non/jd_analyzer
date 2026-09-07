"""
채용 프로세스 소급 파싱 스크립트
사용법:
  export ANTHROPIC_API_KEY='sk-ant-...'
  python3 jd_process.py
결과: jd_analysis.xlsm에 jd_process 시트 추가
"""

import os, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import anthropic

EXCEL_PATH = "/Users/maczoo/Documents/Vault/Attached/jd_analysis.xlsm"
API_KEY    = os.environ.get("ANTHROPIC_API_KEY", "")
MODEL      = "claude-sonnet-4-6"

SYSTEM = """채용공고 원문에서 채용 프로세스 정보를 추출하세요.
반드시 JSON만 반환하고 마크다운 코드블록은 포함하지 마세요.

반환 형식:
{
  "단계": ["서류전형", "코딩테스트", "1차 면접", "2차 면접"],
  "코딩테스트": true 또는 false,
  "과제전형": true 또는 false,
  "명시여부": true 또는 false
}

규칙:
- 채용 프로세스가 명시되어 있지 않으면 명시여부를 false, 단계는 빈 배열로 반환
- 코딩테스트: 코딩테스트/코테/기술 테스트 등이 전형에 포함되면 true
- 과제전형: 과제/사전과제/과제 제출 등이 전형에 포함되면 true"""


def write_process_sheet(wb, results):
    if "jd_process" in wb.sheetnames:
        del wb["jd_process"]
    ws = wb.create_sheet("jd_process")

    HEADER_FILL = PatternFill("solid", fgColor="FF2D2D2D")
    HEADER_FONT = Font(name="맑은 고딕", bold=True, color="FFFFFFFF", size=11)
    TRACK_FILLS = {"A": "FFE8F0FE", "B": "FFFCE8E6", "C": "FFE6F4EA"}

    headers = ["트랙", "공고ID", "코딩테스트", "과제전형", "채용 단계", "프로세스 명시여부"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 50
    ws.column_dimensions['F'].width = 16
    ws.freeze_panes = "A2"

    for r, res in enumerate(results, start=2):
        track = res.get('track', '')
        fill  = PatternFill("solid", fgColor=TRACK_FILLS.get(track, "FFFFFFFF"))
        단계   = " → ".join(res.get('단계') or []) if res.get('단계') else "정보 없음"
        코테   = "✓" if res.get('코딩테스트') else ""
        과제   = "✓" if res.get('과제전형') else ""
        명시   = "명시" if res.get('명시여부') else "미명시"

        ws.cell(row=r, column=1, value=f"트랙 {track}").fill = fill
        ws.cell(row=r, column=2, value=res.get('job_id'))
        ws.cell(row=r, column=3, value=코테).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=4, value=과제).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=5, value=단계)
        ws.cell(row=r, column=6, value=명시).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")


def main():
    if not API_KEY:
        print("❌ ANTHROPIC_API_KEY 환경변수가 설정되지 않았습니다.")
        return

    print(f"📂 파일 열기: {EXCEL_PATH}")
    wb     = openpyxl.load_workbook(EXCEL_PATH, keep_vba=True)
    ws_raw = wb['raw_data']
    client = anthropic.Anthropic(api_key=API_KEY)

    last_row = 1
    for row in ws_raw.iter_rows(min_row=1, max_row=999):
        if any(c.value is not None for c in row):
            last_row = row[0].row

    results = []
    for r in range(2, last_row+1):
        job_id = ws_raw.cell(row=r, column=1).value
        track  = ws_raw.cell(row=r, column=2).value
        raw    = ws_raw.cell(row=r, column=4).value or ''

        print(f"⏳ [{track}] {job_id} 파싱 중...")
        try:
            resp   = client.messages.create(
                model=MODEL, max_tokens=500, system=SYSTEM,
                messages=[{"role": "user", "content": f"채용 프로세스 추출:\n\n{raw[:3000]}"}]
            )
            parsed = json.loads(resp.content[0].text.strip())
            parsed.update({'job_id': job_id, 'track': track})
            results.append(parsed)
            코테 = "O" if parsed.get('코딩테스트') else "-"
            과제 = "O" if parsed.get('과제전형') else "-"
            print(f"   ✅ 코테:{코테} 과제:{과제} | {' → '.join(parsed.get('단계') or [])}")
        except Exception as e:
            print(f"   ❌ 오류: {e}")
            results.append({'job_id': job_id, 'track': track, '명시여부': False, '코딩테스트': False, '과제전형': False, '단계': []})

    print("\n📝 jd_process 시트 작성 중...")
    write_process_sheet(wb, results)
    wb.save(EXCEL_PATH)
    print(f"🎉 완료 — jd_process 시트 저장됨")

    # 트랙별 요약 출력
    print("\n📊 트랙별 요약")
    for track in ['A', 'B', 'C']:
        t_res = [r for r in results if r.get('track') == track]
        total = len(t_res)
        코테_n = sum(1 for r in t_res if r.get('코딩테스트'))
        과제_n = sum(1 for r in t_res if r.get('과제전형'))
        print(f"  트랙 {track} ({total}개): 코딩테스트 {코테_n}개({round(코테_n/total*100)}%) / 과제전형 {과제_n}개({round(과제_n/total*100)}%)")


if __name__ == "__main__":
    main()